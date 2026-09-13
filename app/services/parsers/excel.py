from pathlib import Path

from openpyxl import load_workbook


def parse(path: Path) -> str:
    """逐行读取 xlsx，单元格用 | 分隔；多 sheet 依次拼接。"""
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    finally:
        workbook.close()
