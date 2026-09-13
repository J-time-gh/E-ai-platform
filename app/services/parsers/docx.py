from pathlib import Path

from docx import Document


def parse(path: Path) -> str:
    """提取 docx 的段落与表格文本。"""
    document = Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)
