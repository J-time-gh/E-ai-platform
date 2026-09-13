import subprocess
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from app.services.parsers import parse_document


def test_text_parser_utf8(tmp_path: Path) -> None:
    file = tmp_path / "sample.txt"
    file.write_text("第一行内容\n第二行内容", encoding="utf-8")
    assert "第一行内容" in parse_document(file)


def test_text_parser_gbk(tmp_path: Path) -> None:
    file = tmp_path / "sample_gbk.txt"
    file.write_bytes("GBK 编码的中文内容".encode("gbk"))
    assert "GBK 编码的中文内容" in parse_document(file)


def test_markdown_parser(tmp_path: Path) -> None:
    file = tmp_path / "note.md"
    file.write_text("# 标题\n\n正文段落", encoding="utf-8")
    text = parse_document(file)
    assert "标题" in text
    assert "正文段落" in text


def test_pdf_parser(tmp_path: Path) -> None:
    file = tmp_path / "sample.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Hello PDF content")
    document.save(str(file))
    document.close()
    assert "Hello PDF content" in parse_document(file)


def test_docx_parser(tmp_path: Path) -> None:
    file = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("第一段正文")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "单元格A"
    table.cell(0, 1).text = "单元格B"
    document.save(str(file))

    text = parse_document(file)
    assert "第一段正文" in text
    assert "单元格A | 单元格B" in text


def test_excel_parser(tmp_path: Path) -> None:
    file = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "分数"])
    sheet.append(["小明", 90])
    workbook.save(str(file))

    text = parse_document(file)
    assert "姓名 | 分数" in text
    assert "小明 | 90" in text


def test_unsupported_format_raises(tmp_path: Path) -> None:
    file = tmp_path / "evil.exe"
    file.write_bytes(b"binary")
    with pytest.raises(ValueError, match="不支持"):
        parse_document(file)


def _libreoffice_available() -> bool:
    from app.services.parsers.doc import _find_soffice

    try:
        _find_soffice()
    except RuntimeError:
        return False
    return True


needs_libreoffice = pytest.mark.skipif(not _libreoffice_available(), reason="未安装 LibreOffice")


@needs_libreoffice
def test_doc_parser_via_libreoffice(tmp_path: Path) -> None:
    from app.services.parsers.doc import _find_soffice

    docx_path = tmp_path / "legacy.docx"
    document = Document()
    document.add_paragraph("老格式文档内容")
    document.save(str(docx_path))

    profile = tmp_path / "lo-profile"
    subprocess.run(
        [
            _find_soffice(),
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            "doc",
            "--outdir",
            str(tmp_path),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )

    doc_path = tmp_path / "legacy.doc"
    assert doc_path.exists()
    assert "老格式文档内容" in parse_document(doc_path)


def test_doc_parser_without_libreoffice(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from app.services.parsers import doc as doc_parser

    def no_soffice() -> str:
        raise RuntimeError("未找到 LibreOffice（soffice）")

    monkeypatch.setattr(doc_parser, "_find_soffice", no_soffice)
    with pytest.raises(RuntimeError, match="LibreOffice"):
        doc_parser.parse(tmp_path / "legacy.doc")
